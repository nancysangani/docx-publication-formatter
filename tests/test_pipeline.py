"""
Regression tests for the formatting pipeline.
Run with: pytest tests/
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.pipeline import run_pipeline
from samples.generate_sample import generate, OUT_PATH


def test_pipeline_preserves_content_integrity():
    """The core promise of this project: formatting must never alter content."""
    generate()
    with tempfile.TemporaryDirectory() as tmp:
        output_path = os.path.join(tmp, "formatted.docx")
        result = run_pipeline(OUT_PATH, output_path)
        assert result["integrity"]["passed"] is True
        assert result["integrity"]["mismatched_blocks"] == 0
        assert result["integrity"]["original_checksum"] == result["integrity"]["formatted_checksum"]


def test_pipeline_classifies_all_blocks():
    generate()
    with tempfile.TemporaryDirectory() as tmp:
        output_path = os.path.join(tmp, "formatted.docx")
        result = run_pipeline(OUT_PATH, output_path)
        assert len(result["blocks"]) > 0
        for block in result["blocks"]:
            assert block["label"] in {
                "Title", "Heading", "Subheading", "Body",
                "Caption", "Reference", "List", "Table",
            }


def test_output_file_is_created_and_valid_docx():
    generate()
    with tempfile.TemporaryDirectory() as tmp:
        output_path = os.path.join(tmp, "formatted.docx")
        run_pipeline(OUT_PATH, output_path)
        assert os.path.exists(output_path)
        from docx import Document
        Document(output_path)  # raises if not a valid docx


def test_scales_to_400_plus_pages():
    """The brief's explicit scalability requirement: must handle 400+ page
    manuscripts with acceptable performance and zero content loss."""
    import time
    from samples.generate_large_sample import generate as generate_large

    large_path = generate_large(target_pages=420)
    with tempfile.TemporaryDirectory() as tmp:
        output_path = os.path.join(tmp, "formatted_large.docx")
        t0 = time.time()
        result = run_pipeline(large_path, output_path)
        elapsed = time.time() - t0

        assert result["integrity"]["passed"] is True
        assert result["integrity"]["mismatched_blocks"] == 0
        total_blocks = sum(result["label_counts"].values())
        assert total_blocks > 3000  # sanity check on manuscript size
        assert elapsed < 30, f"Pipeline took {elapsed:.1f}s on a 420-page manuscript (budget: 30s)"

        review_rate = len(result["review_flags"]) / total_blocks
        assert review_rate < 0.05, f"Review rate {review_rate:.1%} exceeds 5% — classifier miscalibrated at scale"


def test_embedded_images_are_preserved():
    """Formatting must never drop embedded content (figures, images) — this
    is what 'preserve all original content' has to mean in practice, not
    just the text. Regression test for a real bug found during development:
    an earlier rebuild-from-scratch formatter silently dropped all images."""
    import io
    from docx import Document as DocxDocument
    from docx.shared import Inches
    from PIL import Image as PILImage

    img = PILImage.new("RGB", (100, 60), color=(120, 40, 40))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    with tempfile.TemporaryDirectory() as tmp:
        input_path = os.path.join(tmp, "with_image.docx")
        output_path = os.path.join(tmp, "formatted_with_image.docx")

        doc = DocxDocument()
        doc.add_paragraph("Chapter 1: Test")
        p = doc.add_paragraph()
        p.add_run().add_picture(buf, width=Inches(1))
        doc.add_paragraph("Fig. 1 A test figure")
        doc.save(input_path)

        original_shapes = len(DocxDocument(input_path).inline_shapes)
        assert original_shapes == 1

        run_pipeline(input_path, output_path)
        formatted_shapes = len(DocxDocument(output_path).inline_shapes)
        assert formatted_shapes == original_shapes, (
            f"Expected {original_shapes} embedded image(s) preserved, found {formatted_shapes}"
        )


if __name__ == "__main__":
    test_pipeline_preserves_content_integrity()
    test_pipeline_classifies_all_blocks()
    test_output_file_is_created_and_valid_docx()
    test_scales_to_400_plus_pages()
    test_embedded_images_are_preserved()
    print("All tests passed.")
